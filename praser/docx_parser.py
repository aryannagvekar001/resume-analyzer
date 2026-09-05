"""Parser for Microsoft Word .docx resumes."""

from __future__ import annotations

from pathlib import Path


class DocxParser:
    supported_suffixes = {".docx"}

    def parse(self, path: str | Path) -> str:
        try:
            from docx import Document
        except ImportError as error:
            raise RuntimeError(
                "python-docx is required. Run: pip install python-docx"
            ) from error

        document = Document(str(path))

        parts = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]
                if cells:
                    parts.append(" | ".join(cells))

        return "\n".join(parts).strip()